"""Document ingestion: parse (+ OCR for images and scanned PDF pages) -> recursive token-aware
chunking -> embed (multilingual MiniLM, local) -> upsert to Pinecone -> mark ready."""

import asyncio
import logging
import re
import time
import uuid
from collections.abc import Callable
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from sqlalchemy import delete, select

from ..config import get_settings
from ..db import SessionLocal
from ..models import Document, DocumentChunk, DocumentEmbedding, utcnow
from . import providers

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """A problem with the file itself; the message is shown to the student."""


@dataclass
class Chunk:
    page: int | None
    text: str


# A PDF page with fewer real characters than this has no usable text layer (a scan or photo).
MIN_PAGE_TEXT_CHARS = 20
_WHITESPACE_RE = re.compile(r"\s+")


@lru_cache(maxsize=4)
def _tesseract_languages(requested: str) -> str:
    """The requested OCR languages that are actually installed (falls back to English)."""
    import pytesseract

    try:
        installed = set(pytesseract.get_languages(config=""))
    except pytesseract.TesseractNotFoundError as exc:
        raise IngestionError("Text recognition (Tesseract OCR) is not installed on the server.") from exc
    languages = [lang for lang in requested.split("+") if lang in installed]
    return "+".join(languages) or "eng"


def ocr_image(image) -> str:
    """OCR a PIL image, after the clean-up that most helps phone photos and screenshots."""
    import pytesseract
    from PIL import ImageOps

    languages = _tesseract_languages(get_settings().ocr_languages)
    image = ImageOps.exif_transpose(image).convert("L")  # respect phone rotation; greyscale
    if image.width < 1000:  # small images: enlarge so characters are big enough to read
        scale = 1000 / max(image.width, 1)
        image = image.resize((int(image.width * scale), int(image.height * scale)))
    image = ImageOps.autocontrast(image)
    return pytesseract.image_to_string(image, lang=languages)


def _ocr_pdf_pages(path: Path, page_numbers: list[int], dpi: int) -> dict[int, str]:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        texts = {}
        for number in page_numbers:
            page = pdf[number - 1]
            try:
                texts[number] = ocr_image(page.render(scale=dpi / 72).to_pil())
            finally:
                page.close()
        return texts
    finally:
        pdf.close()


def extract_text(path: Path) -> tuple[list[tuple[int | None, str]], int]:
    """Returns (page number or None, text) pairs, and how many pages needed OCR."""
    settings = get_settings()
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        pages = [(i, page.extract_text() or "") for i, page in enumerate(PdfReader(str(path)).pages, 1)]
        scanned = [n for n, text in pages if len(_WHITESPACE_RE.sub("", text)) < MIN_PAGE_TEXT_CHARS]
        if not scanned:
            return pages, 0
        if len(scanned) > settings.ocr_max_pages:
            logger.warning("%s: %d scanned pages; reading only the first %d", path.name, len(scanned), settings.ocr_max_pages)
            scanned = scanned[: settings.ocr_max_pages]
        ocr_text = _ocr_pdf_pages(path, scanned, settings.ocr_dpi)
        return [(n, ocr_text.get(n, text)) for n, text in pages], len(scanned)
    if ext == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(None, "\n".join(parts))], 0
    if ext in (".txt", ".md"):
        return [(None, path.read_text(encoding="utf-8", errors="replace"))], 0
    if ext in (".png", ".jpg", ".jpeg"):
        from PIL import Image

        with Image.open(path) as image:
            return [(None, ocr_image(image))], 1
    raise IngestionError(f"Unsupported file type: {ext}")


# Coarsest to finest: paragraphs, lines, sentences, clauses, words, and characters as a last resort.
SEPARATORS = ("\n\n", "\n", ". ", "? ", "! ", "; ", ", ", " ", "")


Length = Callable[[str], int]


def _split_pieces(text: str, size: int, separators: tuple[str, ...], length: Length) -> list[str]:
    """Cut text into pieces of at most `size` (as measured by `length`), using the coarsest
    separator that works. Only pieces that are still too long are split again, more finely.
    """
    if length(text) <= size:
        return [text]
    separator, finer = separators[0], separators[1:]
    if not separator:
        # Last resort for text with no spaces at all. A token is at least one character, so
        # slices of `size` characters are never longer than `size` tokens either.
        return [text[i : i + size] for i in range(0, len(text), size)]
    if separator not in text:
        return _split_pieces(text, size, finer, length)
    parts = text.split(separator)
    pieces: list[str] = []
    for i, part in enumerate(parts):
        # Keep the separator with the text before it, so joined pieces read like the original.
        piece = part + separator if i < len(parts) - 1 else part
        if piece:
            pieces.extend([piece] if length(piece) <= size else _split_pieces(piece, size, finer, length))
    return pieces


def split_text(text: str, size: int, overlap: int, length: Length = len) -> list[str]:
    """Recursive chunking: keep paragraphs (then lines, sentences, …) whole whenever they fit,
    merging neighbours up to `size`. Consecutive chunks share up to `overlap` of whole trailing
    sentences/lines, so context isn't lost at a boundary. Sizes are characters by default, or
    the embedding model's tokens when `length` counts them."""
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    if not text:
        return []
    chunks: list[str] = []
    window: list[tuple[str, int]] = []
    total = 0
    for piece in _split_pieces(text, size, SEPARATORS, length):
        piece_length = length(piece)
        if window and total + piece_length > size:
            chunks.append("".join(p for p, _ in window).strip())
            # Carry whole trailing pieces into the next chunk as overlap.
            while window and (total > overlap or total + piece_length > size):
                total -= window.pop(0)[1]
        window.append((piece, piece_length))
        total += piece_length
    if window:
        chunks.append("".join(p for p, _ in window).strip())
    return [chunk for chunk in chunks if chunk]


def chunk_pages(
    pages: list[tuple[int | None, str]], size: int, overlap: int, length: Length = len
) -> list[Chunk]:
    return [Chunk(page, piece) for page, text in pages for piece in split_text(text, size, overlap, length)]


def chunking_for(embedder) -> tuple[int, int, Length]:
    """Chunk in the embedding model's tokens when it can count them, so no chunk is cut off
    when embedded (the multilingual model reads 128 tokens); otherwise in characters."""
    settings = get_settings()
    count_tokens = getattr(embedder, "count_tokens", None)
    if count_tokens is not None:
        return settings.chunk_tokens, settings.chunk_overlap_tokens, count_tokens
    return settings.chunk_size, settings.chunk_overlap, len


class StepTimer:
    """Records how long each ingestion step takes, for one summary log line per document."""

    def __init__(self) -> None:
        self.started = time.perf_counter()
        self.steps: dict[str, float] = {}
        self._mark = self.started

    def lap(self, step: str) -> None:
        now = time.perf_counter()
        self.steps[step] = self.steps.get(step, 0.0) + (now - self._mark)
        self._mark = now

    def summary(self) -> str:
        parts = [f"{step} {seconds:.1f}s" for step, seconds in self.steps.items()]
        return " | ".join(parts + [f"total {time.perf_counter() - self.started:.1f}s"])


async def ingest_document(document_id: uuid.UUID) -> None:
    settings = get_settings()
    timer = StepTimer()
    async with SessionLocal() as session:
        doc = await session.get(Document, document_id)
        if doc is None:
            return
        doc.status, doc.error, doc.updated_at = "processing", None, utcnow()
        await session.commit()
        timer.lap("start")
        namespace = providers.namespace_for(doc.user_id)
        details = f"{doc.filename!r}, {doc.size_bytes // 1024} KB"
        try:
            if not Path(doc.storage_path).exists():
                raise IngestionError("The uploaded file is no longer on the server (it restarted). Please upload it again.")
            pages, ocr_pages = await asyncio.to_thread(extract_text, Path(doc.storage_path))
            timer.lap("parse+ocr" if ocr_pages else "parse")
            embedder = providers.embedder()
            size, overlap, length = chunking_for(embedder)
            chunks = await asyncio.to_thread(chunk_pages, pages, size, overlap, length)
            timer.lap("chunk")
            details += f", {len(pages)} page(s){f', {ocr_pages} via OCR' if ocr_pages else ''}, {len(chunks)} chunks"
            if not chunks:
                raise IngestionError(
                    "No text could be found, even with text recognition. If it's a photo or scan, "
                    "try a sharper, well-lit image."
                )

            # Re-ingestion: drop previous vectors and rows first.
            old_ids = list(await session.scalars(select(DocumentChunk.id).where(DocumentChunk.document_id == doc.id)))
            if old_ids:
                await asyncio.to_thread(providers.vector_store().delete, namespace, old_ids)
                await session.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
            timer.lap("cleanup")

            vectors = await asyncio.to_thread(embedder.embed_documents, [c.text for c in chunks])
            timer.lap("embed")
            rows = [
                DocumentChunk(id=f"{doc.id.hex}-{i}", document_id=doc.id, chunk_index=i, page=c.page, text=c.text)
                for i, c in enumerate(chunks)
            ]
            items = [
                (
                    row.id,
                    vector,
                    {"document_id": str(doc.id), "chunk_index": row.chunk_index, "page": row.page or 0},
                )
                for row, vector in zip(rows, vectors)
            ]
            await asyncio.to_thread(providers.vector_store().upsert, namespace, items)
            timer.lap("pinecone")
            session.add_all(rows)
            await _record_embedding_model(session, doc.id, settings.embedding_model)
            doc.status, doc.chunk_count, doc.updated_at = "ready", len(rows), utcnow()
            await session.commit()
            timer.lap("database")
            logger.info("Ingested document %s (%s) in %s", doc.id, details, timer.summary())
        except Exception as exc:
            logger.exception("Ingestion failed for document %s (%s) after %s", document_id, details, timer.summary())
            await session.rollback()
            doc = await session.get(Document, document_id, populate_existing=True)
            if doc is not None:
                message = str(exc) if isinstance(exc, IngestionError) else "Processing failed. Try uploading again."
                doc.status, doc.error, doc.updated_at = "failed", message, utcnow()
                await session.commit()


async def _record_embedding_model(session, document_id: uuid.UUID, model: str) -> None:
    await session.execute(delete(DocumentEmbedding).where(DocumentEmbedding.document_id == document_id))
    session.add(DocumentEmbedding(document_id=document_id, model=model))


async def reembed_document(document_id: uuid.UUID) -> None:
    """Bring a ready document's vectors up to the configured embedding model.

    If the uploaded file is still on disk, the document is processed again from scratch (so it
    is also re-chunked for the new model). Otherwise, e.g. on hosts whose disk is wiped on
    restart, the chunk text stored in the database is embedded again: no re-upload needed.
    """
    settings = get_settings()
    timer = StepTimer()
    async with SessionLocal() as session:
        doc = await session.get(Document, document_id)
        if doc is None or doc.status != "ready":
            return
        if Path(doc.storage_path).exists():
            await session.close()
            await ingest_document(document_id)
            return
        try:
            chunks = list(
                await session.scalars(
                    select(DocumentChunk).where(DocumentChunk.document_id == doc.id).order_by(DocumentChunk.chunk_index)
                )
            )
            vectors = await asyncio.to_thread(providers.embedder().embed_documents, [c.text for c in chunks])
            timer.lap("embed")
            items = [
                (c.id, v, {"document_id": str(doc.id), "chunk_index": c.chunk_index, "page": c.page or 0})
                for c, v in zip(chunks, vectors)
            ]
            await asyncio.to_thread(providers.vector_store().upsert, providers.namespace_for(doc.user_id), items)
            timer.lap("pinecone")
            await _record_embedding_model(session, doc.id, settings.embedding_model)
            await session.commit()
            logger.info(
                "Re-embedded document %s (%r, %d chunks) with %s in %s",
                doc.id, doc.filename, len(chunks), settings.embedding_model, timer.summary(),
            )
        except Exception:
            # The document stays usable through keyword search; it is retried at the next startup.
            logger.exception("Re-embedding failed for document %s after %s", document_id, timer.summary())
            await session.rollback()
